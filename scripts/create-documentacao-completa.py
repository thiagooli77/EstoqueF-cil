from pathlib import Path
import shutil

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REFERENCE = Path(r"C:\Users\Thiago77\Downloads\Documentacao_EstoqueFacil.docx")
OUTPUT = ROOT / "Documentacao_EstoqueFacil_Completa.docx"
SCREENSHOTS = Path(r"C:\Users\Thiago77\AppData\Local\Temp\estoquefacil_documentacao_imagens")

BLUE = "2E74B5"
DARK = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = "666666"


def set_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Arial"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{side}"))
        if element is None:
            element = OxmlElement(f"w:{side}")
            tc_mar.append(element)
        element.set(qn("w:w"), str(value))
        element.set(qn("w:type"), "dxa")


def table_geometry(table, widths):
    table.autofit = False
    total = sum(widths)
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    for grid_col, width in zip(table._tbl.tblGrid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def mark_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for cell, label in zip(table.rows[0].cells, headers):
        shade(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = Pt(0)
        set_font(p.add_run(label), size=8.5, bold=True, color=DARK)
    mark_header(table.rows[0])
    for row_number, values in enumerate(rows):
        cells = table.add_row().cells
        for col_number, (cell, value) in enumerate(zip(cells, values)):
            if row_number % 2:
                shade(cell, LIGHT_GRAY)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.05
            if col_number == 0:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_font(p.add_run(value), size=8.5)
    table_geometry(table, widths)
    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True
    set_font(p.add_run(text), size=16 if level == 1 else 13, bold=True, color=BLUE if level == 1 else DARK)
    return p


def add_text(doc, text, lead=None):
    p = doc.add_paragraph(style="Normal")
    if lead and text.startswith(lead):
        set_font(p.add_run(lead), bold=True)
        set_font(p.add_run(text[len(lead):]))
    else:
        set_font(p.add_run(text))
    return p


def add_caption(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(3)
    p.paragraph_format.space_after = Pt(10)
    set_font(p.add_run(text), size=9, italic=True, color=MUTED)


def clear_document_body(document):
    body = document._element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)


def configure(document):
    section = document.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.35)
    section.footer_distance = Inches(0.35)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.15

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 16, 8),
        ("Heading 2", 13, DARK, 11, 5),
    ):
        style = document.styles[name]
        style.font.name = "Arial"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    set_font(header.add_run("Estoque Fácil | Documentação do Projeto"), size=8, color=MUTED)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    set_font(footer.add_run("Documentação técnica e funcional"), size=8, color=MUTED)


def add_screenshot(doc, filename, caption):
    path = SCREENSHOTS / filename
    if not path.exists():
        raise FileNotFoundError(f"Captura não encontrada: {path}")
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(0)
    p.add_run().add_picture(str(path), width=Inches(6.25))
    add_caption(doc, caption)


def build():
    if not REFERENCE.exists():
        raise FileNotFoundError(REFERENCE)
    shutil.copy2(REFERENCE, OUTPUT)
    doc = Document(OUTPUT)
    clear_document_body(doc)
    configure(doc)

    # Capa criada a partir da cópia do exemplo, mantendo sua seção e estilos principais.
    cover_logo = ROOT / "imagens" / "logo-150x100.png"
    cover = doc.add_paragraph()
    cover.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cover.paragraph_format.space_before = Pt(30)
    if cover_logo.exists():
        cover.add_run().add_picture(str(cover_logo), width=Inches(1.45))
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(18)
    title.paragraph_format.space_after = Pt(8)
    set_font(title.add_run("DOCUMENTAÇÃO COMPLETA"), size=25, bold=True, color=DARK)
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(4)
    set_font(subtitle.add_run("Estoque Fácil"), size=18, bold=True, color=BLUE)
    scope = doc.add_paragraph()
    scope.alignment = WD_ALIGN_PARAGRAPH.CENTER
    scope.paragraph_format.space_after = Pt(24)
    set_font(scope.add_run("Sistema de gestão para food trucks e pequenos estabelecimentos"), size=11, italic=True, color=MUTED)
    add_text(doc, "Versão documentada: aplicação web em Node.js, Express e MySQL, com interface HTML/CSS/JavaScript e integração de pagamentos PIX.")
    add_text(doc, "Este documento consolida a visão funcional, arquitetura, banco de dados, rotas, operação, imagens da interface e orientações de uso do projeto.")
    doc.add_page_break()

    add_heading(doc, "RESUMO", 1)
    add_text(doc, "O Estoque Fácil é uma aplicação web voltada à organização de operações de food trucks, lanchonetes e pequenos restaurantes. O sistema reúne cadastro de produtos, ingredientes, receitas, fornecedores, estoque, mesas, pedidos, cozinha, caixa, vendas e pagamentos em uma única solução.")
    add_text(doc, "A aplicação foi organizada em uma arquitetura MVC no back-end e páginas estáticas na pasta views. A API centraliza regras de negócio, persistência no MySQL e integração com a AbacatePay para geração e acompanhamento de cobranças PIX.")

    add_heading(doc, "1. INTRODUÇÃO", 1)
    add_heading(doc, "1.1 Contexto e problema", 2)
    add_text(doc, "Negócios de alimentação de pequeno porte costumam operar com anotações manuais, planilhas isoladas ou controles pouco integrados. Isso aumenta o risco de perdas de estoque, erros de pedido, falta de visão sobre custos e lentidão no fechamento de vendas.")
    add_heading(doc, "1.2 Objetivo do sistema", 2)
    add_text(doc, "O objetivo do Estoque Fácil é dar visibilidade à operação diária: o usuário cadastra itens e insumos, controla entradas e saídas, organiza mesas e pedidos, acompanha a cozinha e registra pagamentos e vendas.")
    add_heading(doc, "1.3 Público-alvo", 2)
    add_text(doc, "O sistema é indicado para proprietários, gerentes, atendentes, cozinheiros e operadores de caixa de food trucks, lanchonetes, pequenos restaurantes e operações móveis de alimentação.")

    add_heading(doc, "2. VISÃO GERAL DAS FUNCIONALIDADES", 1)
    add_table(doc, ["Módulo", "Finalidade", "Principais recursos"], [
        ("Cadastros", "Organizar dados-base da operação.", "Produtos, ingredientes, receitas, categorias, unidades e fornecedores."),
        ("Estoque", "Acompanhar materiais e movimentações.", "Lotes, quantidade, estoque mínimo, recebimentos e status de reposição."),
        ("Atendimento", "Controlar mesas e pedidos.", "Mesas, abertura de pedido, itens, adicionais, observações e status."),
        ("Cozinha", "Acompanhar a produção.", "Quadro de pedidos recebidos, em andamento e concluídos."),
        ("Caixa e vendas", "Registrar fechamento e pagamentos.", "PDV, caixas, vendas, pagamentos e cobranças PIX."),
        ("Gestão", "Apoiar decisão operacional.", "Dashboard, relatórios e consulta dos dados do negócio."),
    ], [1500, 3000, 4860])

    add_heading(doc, "3. TECNOLOGIAS E RECURSOS", 1)
    add_table(doc, ["Camada", "Tecnologias", "Uso no projeto"], [
        ("Servidor", "Node.js e Express", "Inicialização da aplicação, middlewares, rotas e entrega de arquivos estáticos."),
        ("Persistência", "MySQL e mysql2", "Criação do banco, execução de consultas e modelos de dados."),
        ("Interface", "HTML, CSS e JavaScript", "Páginas em views, componentes visuais, formulários e interações."),
        ("API externa", "AbacatePay", "Criação, consulta e simulação de cobranças PIX transparentes e checkouts."),
        ("Segurança", "Token Bearer e hash de senha", "Proteção de rotas autenticadas e armazenamento de credenciais com hash."),
    ], [1500, 2200, 5660])

    add_heading(doc, "4. ARQUITETURA DO SISTEMA", 1)
    add_text(doc, "A aplicação adota uma separação por responsabilidades. As páginas em views apresentam a interface; o Express recebe as requisições; as rotas direcionam cada operação aos controladores; os modelos executam a persistência no MySQL; e serviços especializados concentram integrações externas, como a AbacatePay.")
    add_table(doc, ["Camada", "Arquivos e responsabilidade"], [
        ("Entrada", "server.js inicializa variáveis, banco de dados e servidor HTTP."),
        ("Aplicação", "src/app.js configura parsers, CORS, API, arquivos estáticos e tratamento de erros."),
        ("Rotas", "src/routes organiza autenticação, CRUDs, dashboard e integração de pagamentos."),
        ("Controladores", "src/controllers aplica regras de entrada, respostas HTTP e operações específicas."),
        ("Modelos", "src/models centraliza acesso às tabelas por meio do BaseModel."),
        ("Banco", "src/config/database.js e src/database/mysql_schema.sql configuram MySQL e esquema."),
        ("Interface", "views contém as telas e css/js/imagens fornecem os recursos de apresentação."),
    ], [1800, 7560])

    add_heading(doc, "5. ESTRUTURA DE PASTAS", 1)
    add_table(doc, ["Pasta/arquivo", "Conteúdo"], [
        ("server.js", "Ponto de entrada. Inicia o banco e abre o servidor HTTP."),
        ("src/config", "Configuração de ambiente e conexão com MySQL."),
        ("src/controllers", "Controladores de autenticação, usuários, CRUD, dashboard e AbacatePay."),
        ("src/models", "BaseModel e mapeamento das entidades do banco."),
        ("src/routes", "Rotas de autenticação, CRUD, usuários e pagamentos."),
        ("src/services", "Cliente de comunicação com a API AbacatePay."),
        ("src/middlewares", "Autenticação e tratamento padronizado de erros."),
        ("views", "Páginas HTML do site e módulos administrativos."),
        ("css, js e imagens", "Recursos de estilo, comportamento e mídias visuais."),
        ("scripts", "Inicialização do banco, exportação CSV e automações de apoio."),
    ], [2300, 7060])

    doc.add_page_break()
    add_heading(doc, "6. BANCO DE DADOS", 1)
    add_text(doc, "O banco padrão é estoque_facil. Na inicialização, a aplicação cria o banco quando necessário, aplica o esquema SQL se ainda não houver tabelas e prepara o usuário administrativo de desenvolvimento.")
    add_table(doc, ["Grupo", "Tabelas principais", "Relacionamento funcional"], [
        ("Cadastros", "usuarios, unidades, categorias, fornecedores", "Usuários, unidades de medida, classificação e origem de compras."),
        ("Produtos", "ingredientes, produtos, itens_receita", "Produto final é composto por ingredientes e custos de receita."),
        ("Estoque", "lotes_estoque, movimentacoes_estoque", "Entradas, saídas, lotes, validade e rastreabilidade de movimentos."),
        ("Atendimento", "mesas, pedidos, itens_pedido, adicionais_item_pedido", "Mesa recebe pedido; pedido agrupa itens, observações e adicionais."),
        ("Financeiro", "caixas, vendas, pagamentos, cobrancas_pix", "Pedido pode gerar venda, pagamento e cobrança PIX."),
        ("Compras", "compras, itens_compra", "Registro de compras por fornecedor e itens adquiridos."),
    ], [1600, 3350, 4410])

    add_heading(doc, "7. ROTAS E API", 1)
    add_text(doc, "Todas as rotas de negócio são expostas sob o prefixo /api. A rota /api/health pode ser usada como verificação simples de disponibilidade do servidor.")
    add_table(doc, ["Grupo", "Exemplos de endpoints", "Descrição"], [
        ("Autenticação", "POST /api/auth/login; POST /api/auth/register; GET /api/auth/me", "Login, cadastro e consulta do usuário autenticado."),
        ("CRUD", "GET/POST/PUT/DELETE /api/produtos, /ingredientes, /mesas, /pedidos", "Operações padronizadas sobre entidades do sistema."),
        ("Dashboard", "GET /api/dashboard/summary", "Resumo de informações gerenciais."),
        ("PIX", "POST /api/abacatepay/pix; GET /api/abacatepay/pix/:id", "Criação e consulta de cobrança PIX transparente."),
        ("PIX teste", "POST /api/abacatepay/pix/:id/simular", "Simulação de pagamento durante o modo de teste."),
        ("Checkout", "POST/GET /api/abacatepay/checkouts", "Criação e consulta de checkouts na AbacatePay."),
    ], [1750, 4300, 3310])

    add_heading(doc, "8. AUTENTICAÇÃO E SEGURANÇA", 1)
    add_text(doc, "As rotas protegidas usam o middleware authRequired. O cliente deve enviar o token no cabeçalho Authorization no formato Bearer <token>. Quando o token está ausente ou inválido, a API responde com status 401.")
    add_text(doc, "As configurações de banco de dados e a chave da AbacatePay devem ser mantidas no arquivo .env e não devem ser publicadas no repositório. Antes de publicar a aplicação, as credenciais de desenvolvimento precisam ser substituídas por valores seguros.")

    add_heading(doc, "9. EXECUÇÃO LOCAL", 1)
    add_table(doc, ["Etapa", "Comando/ação", "Resultado esperado"], [
        ("1", "npm install", "Instala as dependências do projeto."),
        ("2", "Configurar o arquivo .env", "Define MySQL e, quando necessário, ABACATEPAY_API_KEY."),
        ("3", "npm run db:init", "Cria ou inicializa a estrutura do banco."),
        ("4", "npm start", "Inicia o servidor; a porta padrão é 3000 e há alternativa para 4000 quando configurada."),
        ("5", "Abrir http://localhost:3000", "Exibe a página inicial servida a partir de views/index.html."),
    ], [650, 3300, 5410])

    doc.add_page_break()
    add_heading(doc, "10. INTERFACE DO PROJETO", 1)
    add_text(doc, "As imagens a seguir foram capturadas da aplicação local durante a elaboração desta documentação. Elas demonstram as telas de acesso, gestão e operação disponíveis no projeto.")
    add_heading(doc, "10.1 Página inicial e planos", 2)
    add_screenshot(doc, "01_inicio.png", "Figura 1 - Página inicial do Estoque Fácil, com funcionalidades, planos e fluxo de pagamento.")
    add_heading(doc, "10.2 Dashboard gerencial", 2)
    add_screenshot(doc, "02_dashboard.png", "Figura 2 - Dashboard gerencial com atalhos para os principais módulos administrativos.")
    add_heading(doc, "10.3 Gestão de estoque", 2)
    add_screenshot(doc, "03_estoque.png", "Figura 3 - Gestão de estoque com busca, quantidade, reposição e status dos ingredientes.")
    add_heading(doc, "10.4 Atendimento de mesas", 2)
    add_screenshot(doc, "04_mesas.png", "Figura 4 - Painel de mesas para organização do atendimento e acompanhamento de status.")
    add_heading(doc, "10.5 Painel da cozinha", 2)
    add_screenshot(doc, "05_cozinha.png", "Figura 5 - Kanban de pedidos da cozinha, separado por recebido, em andamento e concluído.")

    doc.add_page_break()
    add_heading(doc, "11. FLUXO DE FUNCIONAMENTO", 1)
    add_table(doc, ["Etapa", "Fluxo operacional"], [
        ("1. Preparação", "Cadastrar categorias, unidades, fornecedores, ingredientes, produtos e receitas."),
        ("2. Estoque", "Registrar lotes e movimentações; acompanhar quantidades e necessidades de reposição."),
        ("3. Atendimento", "Selecionar mesa, abrir pedido, incluir produtos, adicionais e observações."),
        ("4. Cozinha", "Receber o pedido no painel, avançar o status de preparo e concluir a produção."),
        ("5. Caixa", "Fechar pedido, registrar venda e selecionar a forma de pagamento."),
        ("6. PIX", "Gerar cobrança, exibir QR Code/código de pagamento e consultar ou simular o status."),
        ("7. Gestão", "Consultar dashboard, relatórios e dados para orientar decisões operacionais."),
    ], [1750, 7610])

    add_heading(doc, "12. TESTES E VALIDAÇÃO", 1)
    add_text(doc, "A validação deve cobrir tanto a interface quanto a API. Para cada fluxo, recomenda-se registrar cenário, dados utilizados, resultado esperado, resultado obtido e evidência por captura de tela ou resposta da API.")
    add_table(doc, ["Área", "Cenários mínimos"], [
        ("Autenticação", "Login válido, senha inválida, token ausente e acesso a rota protegida."),
        ("Cadastros", "Criar, consultar, editar e excluir registros de produtos, ingredientes e fornecedores."),
        ("Estoque", "Registrar lote, movimentação, estoque mínimo e apresentação de status."),
        ("Pedidos", "Abrir mesa, adicionar item, incluir observação, atualizar status e fechar pedido."),
        ("Pagamentos", "Gerar PIX, consultar cobrança, copiar código e simular pagamento em teste."),
        ("Interface", "Verificar navegação, responsividade, mensagens de erro e consistência visual."),
    ], [1800, 7560])

    add_heading(doc, "13. CONCLUSÃO", 1)
    add_text(doc, "O Estoque Fácil integra processos que normalmente ficam dispersos em diferentes ferramentas. A combinação de estoque, atendimento, cozinha, caixa e pagamentos cria uma base única para acompanhar a operação de estabelecimentos de alimentação de pequeno porte.")
    add_text(doc, "Como próximos passos, recomenda-se ampliar os relatórios, reforçar permissões por perfil de usuário, automatizar testes, registrar auditoria de ações e validar continuamente a experiência com o público-alvo definido pelo projeto.")

    doc.core_properties.title = "Documentação Completa - Estoque Fácil"
    doc.core_properties.subject = "Documentação técnica e funcional da aplicação Estoque Fácil"
    doc.core_properties.author = "Equipe Estoque Fácil"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build()
