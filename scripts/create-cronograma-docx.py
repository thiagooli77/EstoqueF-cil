from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = Path(__file__).resolve().parents[1] / "Cronograma_Documentacao_15_Dias.docx"

BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"
MUTED = RGBColor(89, 89, 89)


def set_font(run, size=11, bold=False, color=None, italic=False):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    run._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shading = tc_pr.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        tc_pr.append(shading)
    shading.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for name, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{name}"))
        if node is None:
            node = OxmlElement(f"w:{name}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    total = sum(widths)
    table.autofit = False
    tbl_pr = table._tbl.tblPr
    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.append(tbl_w)
    tbl_w.set(qn("w:w"), str(total))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")

    grid = table._tbl.tblGrid
    for grid_col, width in zip(grid.gridCol_lst, widths):
        grid_col.set(qn("w:w"), str(width))

    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
                tc_pr.append(tc_w)
            tc_w.set(qn("w:w"), str(width))
            tc_w.set(qn("w:type"), "dxa")
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    header_cells = table.rows[0].cells
    for cell, value in zip(header_cells, headers):
        set_cell_shading(cell, LIGHT_BLUE)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(value)
        set_font(run, size=9, bold=True, color=DARK_BLUE)
    set_repeat_table_header(table.rows[0])

    for index, values in enumerate(rows):
        cells = table.add_row().cells
        for column, (cell, value) in enumerate(zip(cells, values)):
            if index % 2 == 1:
                set_cell_shading(cell, LIGHT_GRAY)
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            if column == 0:
                paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = paragraph.add_run(value)
            set_font(run, size=9)
    set_table_geometry(table, widths)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def add_heading(doc, text, level=1):
    paragraph = doc.add_paragraph()
    paragraph.style = f"Heading {level}"
    paragraph.paragraph_format.keep_with_next = True
    run = paragraph.add_run(text)
    return paragraph


def add_body(doc, text, bold_prefix=None):
    paragraph = doc.add_paragraph()
    paragraph.style = "Normal"
    if bold_prefix and text.startswith(bold_prefix):
        run = paragraph.add_run(bold_prefix)
        set_font(run, size=11, bold=True)
        run = paragraph.add_run(text[len(bold_prefix):])
        set_font(run, size=11)
    else:
        run = paragraph.add_run(text)
        set_font(run, size=11)
    return paragraph


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
    normal.font.size = Pt(11)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in (
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Calibri")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Calibri")
        style.font.size = Pt(size)
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.bold = True
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    header.paragraph_format.space_after = Pt(0)
    run = header.add_run("Estoque Fácil | Cronograma de Documentação")
    set_font(run, size=8, color="666666")

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.paragraph_format.space_after = Pt(0)
    run = footer.add_run("Documento de planejamento - 15 dias")
    set_font(run, size=8, color="666666")


def build_document():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_before = Pt(42)
    title.paragraph_format.space_after = Pt(8)
    run = title.add_run("CRONOGRAMA DE DOCUMENTAÇÃO")
    set_font(run, size=24, bold=True, color=DARK_BLUE)

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(18)
    run = subtitle.add_run("Estoque Fácil - Pesquisa com Público-Alvo e Documentação do Projeto")
    set_font(run, size=13, color="555555")

    lead = doc.add_paragraph()
    lead.alignment = WD_ALIGN_PARAGRAPH.CENTER
    lead.paragraph_format.space_after = Pt(24)
    run = lead.add_run("Plano de execução em 15 dias, dividido entre quatro integrantes")
    set_font(run, size=10, italic=True, color="666666")

    add_heading(doc, "Objetivo", 1)
    add_body(doc, "Planejar a produção da documentação do sistema Estoque Fácil e validar as necessidades do público-alvo por meio de formulários, entrevistas e testes com usuários reais.")
    add_body(doc, "Importante: as evidências de pesquisa devem conter respostas reais e o uso anônimo das informações deve ser autorizado pelos participantes.", bold_prefix="Importante: ")

    add_heading(doc, "Distribuição da Equipe", 1)
    add_table(
        doc,
        ["Integrante", "Responsabilidade", "Entregas principais"],
        [
            ("Pessoa 1", "Pesquisa e público-alvo", "Personas, formulários, consentimento, entrevistas, respostas, gráficos e relatório da pesquisa."),
            ("Pessoa 2", "Requisitos e processos", "Objetivos, justificativa, requisitos, casos de uso e fluxos de estoque, pedidos, mesas e PDV."),
            ("Pessoa 3", "Documentação técnica", "Arquitetura, banco de dados, APIs, rotas, instalação e estrutura do projeto."),
            ("Pessoa 4", "Qualidade e entrega", "Plano de testes, manual do usuário, capturas de tela, revisão, anexos e PDF final."),
        ],
        [1300, 2500, 5560],
    )

    add_heading(doc, "Ritmo de Trabalho", 1)
    add_body(doc, "Cada integrante atualiza suas entregas ao fim do dia. Nos dias 7, 14 e 15, a equipe faz uma revisão conjunta para verificar evidências, corrigir inconsistências e consolidar o material final.")

    doc.add_page_break()
    add_heading(doc, "Cronograma de 15 Dias", 1)
    add_table(
        doc,
        ["Dia", "Pessoa 1 - Pesquisa", "Pessoa 2 - Requisitos", "Pessoa 3 - Técnico", "Pessoa 4 - Qualidade"],
        [
            ("1", "Definir público-alvo e hipóteses.", "Definir objetivos e justificativa.", "Levantar tecnologias usadas.", "Definir padrão da documentação."),
            ("2", "Criar personas e critérios de seleção.", "Mapear módulos do sistema.", "Mapear estrutura de pastas.", "Criar checklist de qualidade."),
            ("3", "Preparar termo de consentimento.", "Escrever requisitos funcionais iniciais.", "Documentar ambiente e dependências.", "Criar plano de testes."),
            ("4", "Criar formulário de diagnóstico.", "Completar requisitos funcionais.", "Documentar banco de dados.", "Definir evidências de teste."),
            ("5", "Divulgar formulário e registrar canais.", "Levantar requisitos não funcionais.", "Documentar entidades e relações.", "Testar telas principais."),
            ("6", "Coletar respostas e entrevistar usuários.", "Revisar requisitos com achados iniciais.", "Documentar rotas e APIs.", "Registrar falhas encontradas."),
            ("7", "Analisar respostas parciais.", "Criar casos de uso.", "Revisar documentação técnica parcial.", "Revisão conjunta e relatório parcial."),
            ("8", "Produzir gráficos da pesquisa.", "Descrever fluxo de estoque e produtos.", "Documentar arquitetura da aplicação.", "Criar manual do usuário - rascunho."),
            ("9", "Redigir análise das dores dos usuários.", "Descrever fluxo de pedidos e mesas.", "Documentar integração PIX/pagamentos.", "Criar capturas de tela."),
            ("10", "Criar formulário de avaliação do sistema.", "Descrever fluxo de PDV e relatórios.", "Documentar instalação e execução.", "Atualizar roteiro de testes."),
            ("11", "Preparar aplicação dos testes com usuários.", "Revisar requisitos e fluxos.", "Revisar banco, APIs e rotas.", "Testar fluxos completos."),
            ("12", "Coletar avaliações do protótipo.", "Atualizar requisitos conforme retornos.", "Ajustar documentação técnica.", "Registrar resultados de teste."),
            ("13", "Consolidar respostas finais.", "Finalizar casos de uso.", "Finalizar arquitetura e instalação.", "Finalizar manual e evidências."),
            ("14", "Concluir relatório da pesquisa.", "Revisar texto de requisitos.", "Revisar documentação técnica.", "Organizar anexos e formatação."),
            ("15", "Apresentar resultados da pesquisa.", "Validar requisitos finais.", "Validar documentação técnica.", "Consolidar e entregar documento final."),
        ],
        [500, 2215, 2215, 2215, 2215],
    )

    doc.add_page_break()
    add_heading(doc, "Evidências de Pesquisa", 1)
    add_body(doc, "A Pessoa 1 é responsável por manter uma planilha de acompanhamento com data, canal de divulgação, quantidade de respostas, participantes entrevistados e autorização de uso anônimo. Os resultados devem ser anexados ao relatório final.")

    add_heading(doc, "Formulário 1 - Diagnóstico do Público-Alvo", 2)
    add_body(doc, "Aplicar antes ou no início da validação do sistema. Público sugerido: donos de food trucks, lanchonetes, pequenos restaurantes e operadores de caixa.")
    add_table(
        doc,
        ["#", "Pergunta"],
        [
            ("1", "Qual é seu tipo de estabelecimento?"),
            ("2", "Quantas pessoas trabalham no atendimento e no controle de estoque?"),
            ("3", "Como você controla o estoque atualmente?"),
            ("4", "Quais problemas mais acontecem no controle de produtos?"),
            ("5", "Você perde vendas por falta de produto ou erro no estoque?"),
            ("6", "Como registra pedidos e pagamentos hoje?"),
            ("7", "Quais recursos seriam mais úteis: estoque, mesas, pedidos, PDV, relatórios, fornecedores ou cardápio?"),
            ("8", "Usaria um sistema online para centralizar essas tarefas?"),
            ("9", "Quanto estaria disposto a pagar mensalmente?"),
            ("10", "Aceita participar de uma avaliação rápida do sistema?"),
        ],
        [550, 8810],
    )

    add_heading(doc, "Formulário 2 - Avaliação do Sistema", 2)
    add_body(doc, "Aplicar após a demonstração ou uso do protótipo. Registrar as respostas e destacar as melhorias mais recorrentes.")
    add_table(
        doc,
        ["#", "Pergunta"],
        [
            ("1", "Foi fácil entender o sistema?"),
            ("2", "O cadastro de produtos e estoque ficou claro?"),
            ("3", "O fluxo de pedidos e mesas atende sua rotina?"),
            ("4", "O painel de cozinha seria útil?"),
            ("5", "A geração de pagamentos é compreensível?"),
            ("6", "Qual funcionalidade mais ajudaria seu negócio?"),
            ("7", "O que precisa melhorar?"),
            ("8", "De 0 a 10, qual a chance de recomendar o sistema?"),
            ("9", "Autoriza o uso anônimo da sua resposta no relatório acadêmico?"),
        ],
        [550, 8810],
    )

    add_heading(doc, "Checklist de Entrega", 1)
    add_table(
        doc,
        ["Item", "Responsável", "Critério de conclusão"],
        [
            ("Pesquisa e evidências", "Pessoa 1", "Formulários, respostas reais, gráficos, entrevistas e consentimentos organizados."),
            ("Requisitos e fluxos", "Pessoa 2", "Requisitos funcionais e não funcionais, casos de uso e fluxos revisados."),
            ("Documentação técnica", "Pessoa 3", "Arquitetura, banco, APIs, rotas e instalação descritos e conferidos."),
            ("Validação e entrega", "Pessoa 4", "Testes, manual, capturas, anexos, revisão textual e versão final em PDF/Word."),
        ],
        [2000, 1500, 5860],
    )

    doc.core_properties.title = "Cronograma de Documentação - Estoque Fácil"
    doc.core_properties.subject = "Planejamento de 15 dias para documentação e pesquisa com público-alvo"
    doc.core_properties.author = "Equipe Estoque Fácil"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
